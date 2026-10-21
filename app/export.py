"""
Экспорт результата в xlsx по шаблону 12 колонок (+ колонки закупки), включая
вставку миниатюр фото товара и лист с предупреждениями.
"""
import io
from concurrent.futures import ThreadPoolExecutor

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

import currency
import sourcing
from images import _thumb_from_bytes

# Строго по образцу import (22).xlsx — 12 колонок, порядок и заголовки НЕ менять.
OUTPUT_COLUMNS = [
    "Наименование",
    "Китайское наименование товара",
    "Фото товара",
    "Материал",
    "Описание",
    "Ед. измерения",
    "Количество",
    "Вес брутто",
    "Торговая марка",
    "Артикул",
    "Общая закупочная цена $",
    "Закупочная цена за штуку $",
]

SOURCING_COLUMNS = [
    "Ссылка на товар (Китай)", "Поставщик (продавец)",
    "Закупочная цена (Китай)", "Мин. партия (MOQ)",
    "Ссылка на оплату", "Поиск на площадках",
]
SOURCING_W = {13: 40, 14: 28, 15: 20, 16: 14, 17: 24, 18: 55}


# ----------------------------------------------------------- экспорт (по образцу)
HEADER_FILL = PatternFill(start_color="FFCFE2F3", end_color="FFCFE2F3", fill_type="solid")
HEADER_FONT = Font(name="Arial", size=11, bold=False, color="FF000000")
DATA_FONT = Font(name="Arial", size=11)
LINK_FONT = Font(name="Arial", size=11, color="FF0563C1", underline="single")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
# ширины колонок из образца import (22).xlsx
COL_W = {1: 17.71, 2: 19, 3: 13, 4: 13, 5: 13, 6: 14.71, 7: 14.29, 8: 13,
         9: 16.43, 10: 13, 11: 29.57, 12: 37}
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r", "\n")
HEADER_ROW = 2   # строка 1 — номера колонок, строка 2 — заголовки, данные с 3-й
DATA_START = 3


def _embed_images(ws, items):
    """Вставляет картинки СТРОГО В ЯЧЕЙКУ «Фото товара» (C) через two-cell anchor
    (editAs=twoCell) — картинка привязана к обоим углам ячейки, Битрикс видит её
    в нужной клетке и не путает наименования.
    Приоритет: картинка из ИСХОДНОГО инвойса -> найденная в вебе -> заглушка."""
    try:
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
    except Exception:
        return
    targets = [(i + DATA_START, it) for i, it in enumerate(items)]

    def dl(pair):
        r, it = pair
        if it.get("src_image"):          # из файла ИЛИ найдено в вебе (enrich._fetch_web_photos)
            return r, _thumb_from_bytes(it["src_image"])
        return r, None                   # интернет-картинки отключены

    with ThreadPoolExecutor(max_workers=6) as ex:
        results = list(ex.map(dl, targets))
    INSET = 12700  # ~1px отступ от границ ячейки (EMU)
    for r, res in results:
        if not res:
            cell = ws.cell(r, 3)
            if not cell.value:
                cell.value = "Изображение не найдено"
                cell.font = Font(name="Arial", size=9, italic=True, color="999999")
            continue
        buf, (w, h) = res
        try:
            xi = XLImage(buf)
            # C{r} = колонка index 2, строка index r-1 (0-based) .. до D{r+1}
            _from = AnchorMarker(col=2, colOff=INSET, row=r - 1, rowOff=INSET)
            _to = AnchorMarker(col=3, colOff=-INSET, row=r, rowOff=-INSET)
            xi.anchor = TwoCellAnchor(editAs="twoCell", _from=_from, to=_to)
            ws.add_image(xi)
            ws.cell(r, 3).value = None
        except Exception:
            continue


def build_workbook(items, rates, with_sourcing=False, embed_images=True, warnings=None):
    """rates — курсы ЦБ из currency.get_rates(); все суммы пишутся в долларах."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Result"
    columns = OUTPUT_COLUMNS + (SOURCING_COLUMNS if with_sourcing else [])
    ncols = len(columns)

    # строка 1 — номера колонок (как в образце)
    for ci in range(1, ncols + 1):
        c = ws.cell(1, ci, ci)
        c.alignment = CENTER
        c.font = DATA_FONT
    # строка 2 — заголовки (заливка CFE2F3, по центру, без жирного)
    for ci, title in enumerate(columns, 1):
        c = ws.cell(HEADER_ROW, ci, title)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER

    # данные с 3-й строки
    # колонки с одной однозначной ссылкой — делаем кликабельным гиперссылку
    # (а не просто текст URL). «Поиск на площадках» (links_text) сюда
    # намеренно не входит — там 5 РАЗНЫХ ссылок в одной ячейке, у ячейки
    # xlsx может быть только ОДНА цель гиперссылки, нечестно повесить одну
    # на весь текст — оставляем как есть, просто текст со всеми 5 ссылками.
    DESC_COL, CHINA_URL_COL = 5, 13
    for i, it in enumerate(items):
        r = DATA_START + i
        desc = it["description"] or ""
        if it["product_url"]:
            desc = (desc + f"\nСсылка: {it['product_url']}").strip()
        values = [
            it["name_ru"] or it["name_original"] or "",
            it["name_zh"] or "",
            it.get("photo_url", ""),   # заменится картинкой в _embed_images
            it["material"] or "",
            desc,
            it["unit"],
            it["quantity"],
            it["gross_weight"] or "",
            it["brand"] or "",
            it["article"] or it["internal_code"],
            usd(it, "total_price", rates),
            usd(it, "unit_price", rates, ndigits=4),
        ]
        china_url = ""
        if with_sourcing:
            s = it.get("sourcing") or {}
            china_url = s.get("china_url", "")
            values += [
                china_url, s.get("supplier", ""),
                s.get("price_cn", ""), s.get("moq", ""),
                s.get("pay", ""), sourcing.links_text(s.get("links", {})),
            ]
        for ci, v in enumerate(values, 1):
            c = ws.cell(r, ci, v)
            if isinstance(v, str) and v[:1] in _FORMULA_START:
                # текст из инвойса/ИИ не должен стать формулой (=HYPERLINK, DDE…)
                c.data_type = "s"
                c.quotePrefix = True
            if ci == DESC_COL and it["product_url"]:
                c.hyperlink = it["product_url"]
                c.font = LINK_FONT
            elif ci == CHINA_URL_COL and china_url:
                c.hyperlink = china_url
                c.font = LINK_FONT
            else:
                c.font = DATA_FONT
            c.alignment = CENTER
        ws.row_dimensions[r].height = 43.35

    # высоты и ширины по образцу
    ws.row_dimensions[1].height = 14.25
    ws.row_dimensions[HEADER_ROW].height = 42.75
    widths = dict(COL_W)
    if with_sourcing:
        widths.update(SOURCING_W)
    for ci, w in widths.items():
        ws.column_dimensions[get_column_letter(ci)].width = w

    if embed_images:
        _embed_images(ws, items)
    if warnings:
        _add_warnings_sheet(wb, warnings)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _add_warnings_sheet(wb, warnings):
    """Второй лист: что не удалось распознать/перевести. Лист Result остаётся первым."""
    ws = wb.create_sheet("Предупреждения")
    ws.cell(1, 1, "Проверьте вручную — эти данные могут отсутствовать или быть неполными:")
    ws.cell(1, 1).font = Font(name="Arial", size=11, bold=True, color="FFB45309")
    for i, msg in enumerate(warnings, 2):
        c = ws.cell(i, 1, str(msg))
        c.data_type = "s"  # текст с именами файлов не должен стать формулой
        c.font = DATA_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120


def usd(it, key, rates, ndigits=2):
    """Сумма товара в долларах (по умолчанию 2 знака — для цены за штуку передают 4,
    иначе копеечная позиция округляется в $0.00) или "" — если валюта/курс неизвестны:
    пустая ячейка лучше, чем юани или рубли под видом долларов."""
    v = currency.to_usd(it.get(key), it.get("currency"), rates)
    return "" if v is None else round(v, ndigits)
