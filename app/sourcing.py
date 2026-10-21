"""
Модуль закупки на китайских площадках.
1) Deep-ссылки (без ключей) — поиск товара на 1688/Taobao/Tmall/Alibaba/Made-in-China.
2) Адаптер Oxylabs (при наличии OXYLABS_USER/OXYLABS_PASS) — реальный fetch страницы
   1688 через их прокси + best-effort разбор цены/MOQ/поставщика.
Всё аккуратно деградирует до deep-ссылок, если ключей нет или разбор не удался.
"""
import os
import re
import io
import json
import base64
import urllib.request
import urllib.error
from urllib.parse import quote

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

OXY_USER = os.getenv("OXYLABS_USER", "")
OXY_PASS = os.getenv("OXYLABS_PASS", "")
OXY_ENDPOINT = os.getenv("OXYLABS_ENDPOINT", "https://realtime.oxylabs.io/v1/queries")


def _kw(item) -> str:
    """Поисковый запрос: китайское название + артикул."""
    parts = []
    if item.get("name_zh"):
        parts.append(re.sub(r"\s+", " ", item["name_zh"]).strip())
    art = item.get("article") or ""
    # добавляем латинский артикул (хорошо ищется), китайский текст уже есть
    if art and re.search(r"[A-Za-z0-9]", art):
        parts.append(art.split()[0])
    kw = " ".join(parts).strip()
    return kw or item.get("name_ru", "") or art


def deep_links(item) -> dict:
    """Готовые ссылки-поиски на китайских площадках.

    Проверено вручную в браузере 20.10.2026 на реальном запросе (китайское
    название + артикул электродвигателя) — исправлены 2 найденных бага:
    - 1688 ждёт параметр `keywords` в кодировке GBK, а не UTF-8 — с UTF-8
      сервер 1688 отдаёт кракозябры в заголовке и результаты поиска вообще
      не по теме (вместо электродвигателей — случайные украшения и т.п.).
      Кодируем отдельно под GBK; если в запросе есть символы, которых нет в
      GBK (редкость для китайского текста, но бывает), берём что получится
      (errors="ignore") — лучше неполный, но валидный GBK-запрос, чем
      сломанная кодировка.
    - Made-in-China: старый путь `/productdirectory.do?word=...` отдаёт
      "The information you are looking for is not available currently" —
      эндпоинт сам по себе рабочий, но требует параметр `subaction=hunt`
      (подсмотрено из реальной формы поиска на сайте), без него не ищет.
    Taobao/Tmall кодировку (UTF-8) и сам эндпоинт не меняли — они рабочие,
    но требуют быть залогиненным в аккаунт Taobao в браузере, иначе вместо
    результатов показывают "请登录" (пожалуйста, войдите) — это ограничение
    самой площадки, в коде это не обойти."""
    kw = _kw(item)
    q = quote(kw)
    q_gbk = quote(kw.encode("gbk", errors="ignore"))
    links = {
        "1688": f"https://s.1688.com/selloffer/offer_search.htm?keywords={q_gbk}",
        "Taobao": f"https://s.taobao.com/search?q={q}",
        "Tmall": f"https://list.tmall.com/search_product.htm?q={q}",
        "Alibaba": f"https://www.alibaba.com/trade/search?SearchText={q}",
        "Made-in-China": f"https://www.made-in-china.com/productdirectory.do?subaction=hunt&word={q}",
    }
    return links


# --------------------------------------------------------- Oxylabs (опционально)
def oxylabs_available() -> bool:
    return bool(OXY_USER and OXY_PASS)


def _oxy_fetch(url: str, render: bool = False, timeout: int = 60) -> str:
    """Скачивает страницу через Oxylabs Web Scraper API. Возвращает HTML или ''."""
    if not oxylabs_available():
        return ""
    payload = {"source": "universal", "url": url}
    if render:
        payload["render"] = "html"
    data = json.dumps(payload).encode()
    req = urllib.request.Request(OXY_ENDPOINT, data=data, method="POST",
                                 headers={"Content-Type": "application/json"})
    token = base64.b64encode(f"{OXY_USER}:{OXY_PASS}".encode()).decode()
    req.add_header("Authorization", f"Basic {token}")
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            d = json.loads(r.read().decode("utf-8", "ignore"))
        results = d.get("results") or []
        if results:
            return results[0].get("content") or ""
    except Exception:
        return ""
    return ""


_RE_OFFER = re.compile(r"detail\.1688\.com/offer/(\d+)\.html")


def _parse_1688_search(html: str) -> str:
    m = _RE_OFFER.search(html)
    if m:
        return f"https://detail.1688.com/offer/{m.group(1)}.html"
    return ""


def _first(patterns, html):
    for p in patterns:
        m = re.search(p, html)
        if m:
            return m.group(1).strip()
    return ""


def _parse_1688_product(html: str) -> dict:
    price = _first([
        r'"referencePrice"\s*:\s*"?([\d.]+)',
        r'"price"\s*:\s*"?([\d.]+)',
        r'¥\s*([\d.]+)',
    ], html)
    moq = _first([
        r'起订量[^\d]{0,6}(\d+)',
        r'"minOrderQuantity"\s*:\s*"?(\d+)',
        r'最小起订[^\d]{0,6}(\d+)',
    ], html)
    supplier = _first([
        r'"companyName"\s*:\s*"([^"]{2,80})"',
        r'"sellerLoginId"\s*:\s*"([^"]{2,80})"',
        r'公司名称[:：]\s*([^<\n]{2,80})',
    ], html)
    return {"price": price, "moq": moq, "supplier": supplier}


def source_via_oxylabs(item) -> dict:
    """Реальный поиск на 1688 через Oxylabs. Возвращает частично заполненный dict."""
    out = {"china_url": "", "supplier": "", "price_cn": "", "moq": ""}
    if not oxylabs_available():
        return out
    kw = _kw(item)
    if not kw:
        return out
    # GBK, не UTF-8 — см. комментарий в deep_links() про тот же параметр
    search_url = f"https://s.1688.com/selloffer/offer_search.htm?keywords={quote(kw.encode('gbk', errors='ignore'))}"
    html = _oxy_fetch(search_url, render=True)
    offer = _parse_1688_search(html)
    if not offer:
        return out
    out["china_url"] = offer
    phtml = _oxy_fetch(offer, render=True)
    if phtml:
        p = _parse_1688_product(phtml)
        if p["price"]:
            out["price_cn"] = f"¥{p['price']}"
        out["moq"] = p["moq"]
        out["supplier"] = p["supplier"]
    return out


def source_item(item) -> dict:
    """Единая точка: deep-ссылки всегда + Oxylabs-данные, если доступны."""
    res = {
        "china_url": "", "supplier": "", "price_cn": "", "moq": "",
        "pay": "",  # QR/оплата — статически недоступно, оставляем пустым
        "links": deep_links(item),
    }
    if oxylabs_available():
        try:
            res.update({k: v for k, v in source_via_oxylabs(item).items() if v})
        except Exception:
            pass
    # если карточку 1688 не нашли — ссылка на товар = deep-поиск 1688
    if not res["china_url"]:
        res["china_url"] = res["links"].get("1688", "")
    return res


def links_text(links: dict) -> str:
    return "\n".join(f"{k}: {u}" for k, u in links.items())


# ---------------------------------------------------- шаблон ТЗ на закупку
# Колонки подобраны под то, что реально использует поиск на площадках
# (_kw() выше: китайское название + артикул — главные поля для качества
# поиска) и под extract.py (та же ИИ-модель, что разбирает инвойсы,
# разбирает и этот файл — колонки похожи на инвойсные специально, чтобы
# не заводить отдельный формат/парсер). Цена необязательна — для закупки
# обычно не известна заранее, в отличие от инвойса.
TZ_COLUMNS = [
    "Наименование (что ищем)",
    "Китайское название 中文名称 (если знаете — сильно повышает точность поиска)",
    "Артикул / модель",
    "Бренд (если нужен конкретный, иначе оставьте пустым — ищем любой)",
    "Количество",
    "Ед. изм.",
    "Характеристики / спецификация (материал, размер, мощность, напряжение и т.п.)",
    "Желаемая цена за шт, $ (необязательно)",
]
TZ_EXAMPLE_ROW = [
    "Асинхронный электродвигатель с тормозом",
    "三相异步电动机(带制动器)",
    "YEJ2-132M-4",
    "WNM",
    10,
    "шт",
    "4 кВт, 380В, IP54, чугунный корпус",
    45,
]
TZ_INSTRUCTIONS = [
    "Как заполнять этот файл, чтобы программа нашла нужные товары",
    "",
    "1. Одна строка = один товар, который нужно найти и закупить в Китае.",
    "2. Заполняйте столько строк, сколько нужно товаров — остальные листы/структуру не трогайте.",
    "3. ВАЖНЕЕ ВСЕГО для качества поиска — китайское название и/или артикул/модель.",
    "   Если это известный китайский бренд/модель (как в примере WNM YEJ2-132M-4) — "
    "почти наверняка найдётся точное совпадение.",
    "4. Если китайского названия не знаете — хотя бы точно укажите тип товара и "
    "ключевые характеристики (материал, размер, мощность, напряжение и т.п.) в "
    "столбце «Характеристики» — по ним тоже можно искать, но менее точно.",
    "5. Бренд указывайте, только если нужен именно он — если подходит любой "
    "производитель (OEM), оставьте пустым.",
    "6. Цена — необязательна, нужна только если у вас есть бюджет/ориентир.",
    "7. После заполнения загрузите этот файл во вкладке «Закупка (Китай)», "
    "как обычный инвойс — программа сама найдёт строки с товарами.",
]


def build_tz_template() -> bytes:
    """Шаблон для заполнения технического задания на закупку — скачивается
    кнопкой «Шаблон ТЗ» во вкладке «Закупка (Китай)»."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ТЗ"
    header_fill = PatternFill(start_color="FFCFE2F3", end_color="FFCFE2F3", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True)
    wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for ci, title in enumerate(TZ_COLUMNS, 1):
        c = ws.cell(1, ci, title)
        c.fill = header_fill
        c.font = header_font
        c.alignment = wrap
    for ci, v in enumerate(TZ_EXAMPLE_ROW, 1):
        c = ws.cell(2, ci, v)
        c.font = Font(name="Arial", size=11, italic=True, color="999999")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 60
    widths = [28, 32, 18, 20, 12, 10, 36, 20]
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    ws2 = wb.create_sheet("Инструкция")
    ws2.column_dimensions["A"].width = 100
    for ri, line in enumerate(TZ_INSTRUCTIONS, 1):
        c = ws2.cell(ri, 1, line)
        c.font = Font(name="Arial", size=12, bold=(ri == 1))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws2.row_dimensions[ri].height = 18 if line else 8

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
