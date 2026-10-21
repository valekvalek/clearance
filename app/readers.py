"""
Чтение исходных файлов инвойса: xlsx/xls -> сетка с координатами ячеек,
PDF -> текстовые листы (постранично) + страницы-сканы для vision.
"""
import io
import os

from openpyxl.utils import get_column_letter
import openpyxl

MAX_COLS = 40    # широкие инвойсы: веса/итоги часто в правых колонках
MAX_ROWS = 3000  # до ~1500+ позиций на лист (крупные инвойсы/пакинги)
MAX_VISION_PAGES = int(os.getenv("MAX_VISION_PAGES", "8"))


# ---------------------------------------------------------------- чтение xlsx/xls
def _grid_from_rows(row_iter, info=None):
    """Из итератора строк (списки значений) -> текст сетки с координатами.
    info (dict) получает truncated=True, если лист длиннее MAX_ROWS."""
    lines = []
    for ri, row in enumerate(row_iter, start=1):
        if ri > MAX_ROWS:
            if info is not None:
                info["truncated"] = True
            break
        cells = []
        has = False
        for ci, v in enumerate(row):
            if ci >= MAX_COLS:
                break
            if v is None or v == "":
                cells.append("")
            else:
                s = str(v).replace("\n", " ").replace("\t", " ").strip()
                cells.append(s[:60])
                if s:
                    has = True
        if has:
            lines.append(
                f"R{ri}: " + " | ".join(
                    f"{get_column_letter(i + 1)}={c}" for i, c in enumerate(cells) if c
                )
            )
    return "\n".join(lines)


def read_sheets(data: bytes):
    """Читает книгу из bytes (.xlsx или .xls) -> [{'name', 'grid'}]."""
    # .xls (OLE2) начинается с D0CF11E0; .xlsx (zip) — с PK
    if data[:4] == b"\xd0\xcf\x11\xe0":
        return _read_xls(data)
    sheets = []
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    for ws in wb.worksheets:
        sh = {"name": ws.title}
        sh["grid"] = _grid_from_rows(ws.iter_rows(values_only=True), sh)
        sheets.append(sh)
    wb.close()
    return sheets


def _read_xls(data: bytes):
    import xlrd
    wb = xlrd.open_workbook(file_contents=data)
    sheets = []
    for sh in wb.sheets():
        rows = ([sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows))
        info = {"name": sh.name}
        info["grid"] = _grid_from_rows(rows, info)
        sheets.append(info)
    return sheets


# ------------------------------------------------------------- PDF-инвойсы
def _read_pdf(data: bytes):
    """PDF -> (sheets, vision_pages, n_scan_pages). Режим выбирается по каждой странице
    отдельно: страница с достаточным текстом попадает в sheets, страница-скан рендерится в
    PNG для vision (не больше MAX_VISION_PAGES штук на файл, остальные только считаются)."""
    import fitz  # PyMuPDF
    doc = fitz.open(stream=data, filetype="pdf")
    sheets, vision_pages = [], []
    for pi in range(len(doc)):
        text = doc[pi].get_text()
        if len(text.strip()) >= 150:
            lines = [f"R{i + 1}: {ln.strip()}"
                     for i, ln in enumerate(text.splitlines()) if ln.strip()][:MAX_ROWS]
            sheets.append({"name": f"PDF стр.{pi + 1}", "grid": "\n".join(lines)})
        elif len(vision_pages) < MAX_VISION_PAGES:
            vision_pages.append(doc[pi].get_pixmap(dpi=140).tobytes("png"))
    n_scan_pages = len(doc) - len(sheets)
    return sheets, vision_pages, n_scan_pages
