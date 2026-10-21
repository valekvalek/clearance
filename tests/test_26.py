"""Шаг 2.6: цена за штуку — 4 знака после запятой (копеечные позиции не дают $0.00
в столбце «Закупочная цена за штуку $»). «Общая закупочная цена $» по-прежнему
в 2 знаках."""
import io, os, sys, tempfile, warnings
warnings.filterwarnings("ignore")
os.environ.setdefault("DATA_DIR", tempfile.mkdtemp())
sys.path.insert(0, sys.argv[1])
import openpyxl, export

ok = True
def check(name, cond):
    global ok; ok &= bool(cond); print(("✓ " if cond else "✗ ") + name)

RATES = {"usd_per": {"USD": 1.0, "CNY": 1 / 7}, "date": "01.01.2026", "source": "ЦБ РФ", "stale": False}

print("— usd(): ndigits по умолчанию 2, явно можно 4 —")
cheap = {"unit_price": 0.001, "total_price": 0.1, "currency": "USD"}
check(f"по умолчанию (для «Общей») копейки округляются до 2 знаков: {export.usd(cheap, 'unit_price', RATES)}",
      export.usd(cheap, "unit_price", RATES) == 0.0)
check(f"с ndigits=4 (для «за штуку») копейки видны: {export.usd(cheap, 'unit_price', RATES, ndigits=4)}",
      export.usd(cheap, "unit_price", RATES, ndigits=4) == 0.001)
check("«Общая закупочная цена» — по-прежнему 2 знака",
      export.usd(cheap, "total_price", RATES) == 0.1)

print("— build_workbook: копеечная позиция видна в колонке «за штуку», не $0.00 —")
def item(name, art, qty, unit_price, currency="USD"):
    return {"name_ru": name, "name_original": name, "name_zh": "", "material": "",
            "description": "", "unit": "шт", "quantity": qty, "gross_weight": 0,
            "brand": "", "article": art, "internal_code": "", "unit_price": unit_price,
            "total_price": round(unit_price * qty, 6), "currency": currency,
            "product_url": "", "photo_url": "", "sourcing": None}

items = [
    item("грошовая деталь", "A1", 1000, 0.0035),   # $0.0035/шт, при 2 знаках -> $0.00
    item("обычный товар", "A2", 2, 9.995),          # $9.995 -> округление до 4 знаков ожидаемо
]
data = export.build_workbook(items, 0.0, with_sourcing=False, embed_images=False)
ws = openpyxl.load_workbook(io.BytesIO(data)).active
unit_col, total_col = 12, 11  # «Закупочная цена за штуку $», «Общая закупочная цена $»
u1, t1 = ws.cell(3, unit_col).value, ws.cell(3, total_col).value
u2, t2 = ws.cell(4, unit_col).value, ws.cell(4, total_col).value
check(f"грошовая деталь: цена за штуку не ноль — {u1}", u1 == 0.0035)
check(f"грошовая деталь: общая сумма в 2 знаках — {t1}", t1 == round(0.0035 * 1000, 2))
check(f"обычный товар: цена за штуку в 4 знаках — {u2}", u2 == round(9.995, 4))
check(f"обычный товар: общая сумма в 2 знаках — {t2}", t2 == round(9.995 * 2, 2))

print("\nИТОГ:", "ВСЁ ОК" if ok else "ЕСТЬ ОШИБКИ")
