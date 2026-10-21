import io, re, sys, zipfile, warnings; warnings.filterwarnings("ignore")
sys.path.insert(0, sys.argv[1])
import openpyxl, export
TRAPS = ['=1+1', '=HYPERLINK("http://evil.example","Нажмите")', "+cmd|'/c calc'!A1",
         '@SUM(1+1)', '-2+3', '\t=1+1', '-Ручка шариковая', 'Обычный товар']
def item(name, art, qty, price):
    return {"name_ru": name, "name_original": name, "name_zh": "=1+2", "material": "@x",
            "description": "+описание", "unit": "шт", "quantity": qty, "gross_weight": 1.5,
            "brand": "-OEM", "article": art, "internal_code": "", "total_price": price,
            "unit_price": price / qty, "currency": "USD", "product_url": "", "photo_url": "",
            "sourcing": {"china_url": "=A1", "supplier": "+7 999", "price_cn": "¥1",
                         "moq": "-", "pay": "", "links": {}}}
items = [item(t, t, 2, 10.0) for t in TRAPS]
data = export.build_workbook(items, 0.14, with_sourcing=True, embed_images=False)
xml = zipfile.ZipFile(io.BytesIO(data)).read("xl/worksheets/sheet1.xml").decode()
ws = openpyxl.load_workbook(io.BytesIO(data)).active
formulas = [(c.coordinate, c.value) for row in ws.iter_rows() for c in row if c.data_type == "f"]
print(f"формул в файле: {len(formulas)} {formulas[:4]}{' …' if len(formulas) > 4 else ''}")
print(f"тегов <f> в XML листа: {len(re.findall(r'<f[ >]', xml))}")
if len(sys.argv) > 2:
    ok = not formulas
    texts = [ws.cell(3 + i, 1).value for i in range(len(TRAPS))]
    same = texts == TRAPS; ok &= same
    print(("✓ " if same else "✗ ") + "текст сохранён дословно:", texts[:3], "…")
    nums = [(ws.cell(3, 7).value, ws.cell(3, 8).value, ws.cell(3, 11).value)]
    isnum = all(isinstance(x, (int, float)) for x in nums[0]); ok &= isnum
    print(("✓ " if isnum else "✗ ") + f"кол-во/вес/сумма остались числами: {nums[0]}")
    qp = ws.cell(3, 1).quotePrefix and not ws.cell(10, 1).quotePrefix; ok &= qp
    print(("✓ " if qp else "✗ ") + "quotePrefix только у «опасных» ячеек (обычный товар без него)")
    print("ИТОГ:", "ВСЁ ОК" if ok and not formulas else "ЕСТЬ ОШИБКИ")
