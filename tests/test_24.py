"""Шаг 2.4: все суммы в долларах по курсу ЦБ РФ."""
import io, json, os, re, sys, tempfile, time, warnings
warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__))
D = tempfile.mkdtemp()
os.environ.update(DATA_DIR=D, ADMIN_USER="boss", ADMIN_PASSWORD="bosspass123")
sys.path.insert(0, sys.argv[1])
import openpyxl
from fastapi.testclient import TestClient
import app as appmod, currency, gemini

XML = open(os.path.join(HERE, "cbr.xml"), "rb").read()        # настоящий ответ cbr.ru на 25.09.2026
JSN = open(os.path.join(HERE, "mirror.json"), "rb").read()    # настоящий ответ зеркала
USD_RUB, EUR_RUB, CNY_RUB, JPY_RUB = 84.9057, 96.8859, 12.65, 53.7514 / 100

ok = True
def check(name, cond):
    global ok; ok &= bool(cond); print(("✓ " if cond else "✗ ") + name)
def near(a, b, eps=1e-6):
    return a is not None and abs(a - b) < eps

print("— написания валют —")
cases = {"RMB": "CNY", "CNY": "CNY", "RMB/CNY": "CNY", "RMB¥": "CNY", "CNY ¥": "CNY", "¥": "CNY",
         "元": "CNY", "人民币": "CNY", "yuan": "CNY", "Юань": "CNY", "": "CNY", "RMB/CNY/USD": "CNY",
         "USD": "USD", "US$": "USD", "$": "USD", "美元": "USD", "US Dollar": "USD",
         "EUR": "EUR", "€": "EUR", "euro": "EUR", "евро": "EUR",
         "RUB": "RUB", "RUR": "RUB", "руб.": "RUB", "₽": "RUB", "рублей": "RUB",
         "HKD": "HKD", "HK$": "HKD", "港币": "HKD", "A$": "AUD", "JPY": "JPY", "日元": "JPY", "£": "GBP",
         "USD/EUR": None, "N/A": None, "???": None}
bad = {k: (currency.normalize_currency(k), v) for k, v in cases.items() if currency.normalize_currency(k) != v}
check(f"{len(cases) - len(bad)}/{len(cases)} написаний распознаны верно" + (f", ошибки: {bad}" if bad else ""), not bad)

print("— курсы ЦБ (настоящие данные на 25.09.2026) —")
calls = []
def fake_get(url, data={"xml": XML, "json": JSN}):
    calls.append(url)
    if url == currency.CBR_XML_URL:
        if isinstance(data["xml"], Exception): raise data["xml"]
        return data["xml"]
    if isinstance(data["json"], Exception): raise data["json"]
    return data["json"]
currency._get = fake_get
def fresh(xml=XML, jsn=JSN):
    currency._CACHE.update(rates=None, at=0); calls.clear()
    fake_get.__defaults__[0].update(xml=xml, json=jsn)
    return currency.get_rates()

r = fresh()
up = r["usd_per"]
check(f"источник {r['source']}, дата {r['date']}, валют {len(up)}", r["source"] == "ЦБ РФ" and r["date"] == "25.09.2026" and len(up) > 40)
check(f"1 USD = {1/up['CNY']:.4f} CNY (= {USD_RUB}/{CNY_RUB})", near(1 / up["CNY"], USD_RUB / CNY_RUB))
check(f"1 EUR = {up['EUR']:.4f} USD (= {EUR_RUB}/{USD_RUB})", near(up["EUR"], EUR_RUB / USD_RUB))
check(f"1 RUB = {up['RUB']:.6f} USD (= 1/{USD_RUB})", near(up["RUB"], 1 / USD_RUB))
check(f"JPY с номиналом 100: 1 JPY = {up['JPY']:.6f} USD", near(up["JPY"], JPY_RUB / USD_RUB))
check("1000 ₽ → $%.2f, 100 € → $%.2f, 671 ¥ → $%.2f" % (currency.to_usd(1000, "RUB", r), currency.to_usd(100, "EUR", r), currency.to_usd(671, "CNY", r)),
      near(currency.to_usd(1000, "RUB", r), 1000 / USD_RUB) and near(currency.to_usd(100, "EUR", r), 100 * EUR_RUB / USD_RUB))
print("  строка курса в интерфейсе:", currency.rate_display(r, ["CNY", "EUR", "RUB", "USD"]))
currency.get_rates()
check("повторный запрос в течение часа — из кеша, без сети", calls == [currency.CBR_XML_URL])
r2 = fresh(xml=OSError("cbr.ru недоступен"))
check(f"cbr.ru упал → зеркало: {r2['source']}, та же цена юаня",
      "cbr-xml-daily" in r2["source"] and near(1 / r2["usd_per"]["CNY"], 1 / up["CNY"], 1e-3))
insane = XML.replace(b"<CharCode>CNY</CharCode><Nominal>1</Nominal><Name>\xde\xe0\xed\xfc</Name><Value>12,6500",
                     b"<CharCode>CNY</CharCode><Nominal>1</Nominal><Name>\xde\xe0\xed\xfc</Name><Value>1,2650")
check("в XML подменён курс юаня (1 USD = 67 CNY)", insane != XML)
r3 = fresh(xml=insane)
check("неправдоподобный курс отбракован → взято зеркало", "cbr-xml-daily" in r3["source"])
r4 = fresh(xml=OSError("нет сети"), jsn=OSError("нет сети"))
check(f"сети нет → сохранённый с диска курс на {r4['date']}, помечен stale", r4["stale"] and r4["usd_per"])
os.remove(currency.RATES_FILE)
r5 = fresh(xml=OSError("нет сети"), jsn=OSError("нет сети"))
check("сети нет и на диске пусто → курсов нет, но доллары остаются долларами",
      r5["usd_per"] == {} and currency.to_usd(5, "USD", r5) == 5 and currency.to_usd(5, "CNY", r5) is None)
fresh()  # вернуть нормальные курсы для прогона задачи

print("— полный прогон: инвойс в разных валютах —")
ROWS = [("item1", 10, 671, "RMB/CNY"), ("item2", 1, 100, "€"), ("item3", 2, 1000, "руб."),
        ("item4", 3, 50, "US$"), ("item5", 1, 999, "XYZ"), ("item6", 1, 5, "???"), ("item7", 4, 20, "")]
def fake_chat(messages, **kw):
    t = messages[-1]["content"]
    if "Разбери инвойс" in t:
        return json.dumps([{"name_original": n, "quantity": q, "unit_price": p / q, "total_price": p,
                            "currency": c, "row": i + 2} for i, (n, q, p, c) in enumerate(ROWS)])
    return json.dumps([{"name_ru": "RU " + n, "brand": "B", "material": "m"} for n in re.findall(r"наименование: (\S+)", t)])
gemini.chat, gemini.API_KEY = fake_chat, "test"
wb = openpyxl.Workbook(); wb.active.append(["x"]); b = io.BytesIO(); wb.save(b)
c = TestClient(appmod.app, base_url="https://testserver")
c.post("/api/login", json={"username": "boss", "password": "bosspass123"})
jid = c.post("/api/process", files=[("files", ("inv.xlsx", b.getvalue()))]).json()["job_id"]
for _ in range(50):
    st = c.get(f"/api/status/{jid}").json()
    if st["status"] in ("done", "error"): break
    time.sleep(0.1)
print("  курс в интерфейсе:", st["rate"])
for w in st["warnings"]:
    print("  ⚠", w)
ws = openpyxl.load_workbook(io.BytesIO(c.get(f"/api/result/{jid}").content))["Result"]
got = {ws.cell(r, 1).value.replace("RU ", ""): (ws.cell(r, 11).value, ws.cell(r, 12).value) for r in range(3, 3 + len(ROWS))}
exp = {"item1": 671 * up["CNY"], "item2": 100 * up["EUR"], "item3": 1000 * up["RUB"], "item4": 50.0, "item7": 20 * up["CNY"]}
for n, v in exp.items():
    check(f"{n}: {dict((r[0], r[3]) for r in ROWS)[n] or '(не указана)'} → «Общая $» = {got[n][0]} (ожидалось {v:.2f})", got[n][0] == round(v, 2))
check(f"item3: цена за шт. тоже в $, округлена до 4 знаков (шаг 2.6) — {got['item3'][1]} (= 500 ₽)",
      got["item3"][1] == round(500 * up["RUB"], 4))
check(f"item5 (XYZ) и item6 (???): ячейки $ пустые, а не «999» под видом долларов: {got['item5']}, {got['item6']}",
      got["item5"] == (None, None) and got["item6"] == (None, None))
check("предупреждение: у ЦБ нет курса XYZ (товар № 5)", any("№ 5" in w and "XYZ" in w for w in st["warnings"]))
check("предупреждение: валюта «???» не распознана (товар № 6)", any("№ 6" in w and "«???»" in w for w in st["warnings"]))
pv = c.get(f"/api/preview/{jid}").json()["preview"]
check("предпросмотр: те же суммы, у неизвестных — пусто",
      [p["total_usd"] for p in pv] == [got[f"item{i}"][0] if got[f"item{i}"][0] is not None else "" for i in range(1, 8)])
check("строка курса перечисляет CNY, EUR, RUB по ЦБ на дату", all(x in st["rate"] for x in ("ЦБ РФ на 25.09.2026", "CNY", "EUR", "RUB")))
print("\nИТОГ:", "ВСЁ ОК" if ok else "ЕСТЬ ОШИБКИ")
