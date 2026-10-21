"""Шаг 2.5: числа из инвойса — валюта/тысячи/десятые в _num, досчёт total ↔ unit_price,
предупреждение о нераспознанной цене. ИИ заменён заглушкой."""
import io, json, os, re, sys, tempfile, time, warnings
warnings.filterwarnings("ignore")
os.environ.update(DATA_DIR=tempfile.mkdtemp(), ADMIN_USER="boss", ADMIN_PASSWORD="bosspass123")
sys.path.insert(0, sys.argv[1])
import openpyxl
from fastapi.testclient import TestClient
import app as appmod, currency, extract, gemini

ok = True
def check(name, cond):
    global ok; ok &= bool(cond); print(("✓ " if cond else "✗ ") + name)

print("— _num/_parse_amount: 20+ форматов —")
cases = [
    ("¥12.5", 12.5, "CNY"), ("12.5 USD", 12.5, "USD"), ("€5", 5.0, "EUR"),
    ("US$1,234.56", 1234.56, "USD"), ("HK$100", 100.0, "HKD"),
    ("12,5", 12.5, None), ("1,234", 1234.0, None), ("1.234,56", 1234.56, None),
    ("1,234.56", 1234.56, None), ("1 234,56", 1234.56, None),
    ("1 234", 1234.0, None), ("100", 100.0, None), ("0", 0.0, None),
    ("-15.5", -15.5, None), ("1.234.567", 1234567.0, None),
    ("1,234,567", 1234567.0, None), ("  42  ", 42.0, None), ("¥ 5", 5.0, "CNY"),
    ("人民币12.5", 12.5, "CNY"), ("руб. 100,50", 100.5, "RUB"), ("", 0.0, None),
    (None, 0.0, None), ("0.99", 0.99, None),
]
bad = []
for raw, exp_val, exp_cur in cases:
    val, cur, _ = extract._parse_amount(raw)
    if abs(val - exp_val) > 1e-6 or (exp_cur is not None and cur != exp_cur):
        bad.append((raw, val, cur))
check(f"{len(cases) - len(bad)}/{len(cases)} форматов разобраны верно" + (f", ошибки: {bad}" if bad else ""), not bad)

print("— нераспознанная цена —")
val, cur, recognized = extract._parse_amount("по запросу")
check("«по запросу»: 0, не распознано", val == 0.0 and recognized is False)
val, cur, recognized = extract._parse_amount("0")
check("«0»: легитимный ноль, распознано", val == 0.0 and recognized is True)

print("— досчёт total ↔ unit_price и валюта из цены —")
data = [
    {"name_original": "A", "quantity": "10", "unit_price": "¥12.5", "total_price": ""},
    {"name_original": "B", "quantity": "4", "unit_price": "", "total_price": "¥100", "currency": ""},
    {"name_original": "C", "quantity": "2", "unit_price": "по запросу", "total_price": ""},
    {"name_original": "D", "quantity": "3", "unit_price": "5", "total_price": "15", "currency": "USD"},
    {"name_original": "E", "quantity": "1", "unit_price": "0", "total_price": "0", "currency": "USD"},
]
items = extract._items_from_json(data)
A, B, C, D, E = items
check(f"A: total досчитан из qty×unit_price = {A['total_price']}", A["total_price"] == 125.0)
check(f"B: unit_price досчитан из total/qty = {B['unit_price']}", B["unit_price"] == 25.0)
check(f"A/B: валюта взята из символа цены (поле валюты пусто) = {A['currency']}, {B['currency']}",
      A["currency"] == "CNY" and B["currency"] == "CNY")
check(f"D: явная валюта из поля не перебита = {D['currency']}", D["currency"] == "USD")
check("C: цена не распознана, отмечена для отчёта", C.get("price_unrecognized") == "по запросу")
check("E: легитимный 0 не помечается как нераспознанный", E.get("price_unrecognized") is None)
check("A/B/D/E: не помечены", all(it.get("price_unrecognized") is None for it in (A, B, D, E)))

print("— полный прогон через API: цены строками, разные форматы —")
currency.get_rates = lambda: {"usd_per": {"USD": 1.0, "CNY": 1 / 7}, "date": "01.01.2026", "source": "ЦБ РФ", "stale": False}
ROWS = [
    ("item1", "10", "¥12.5", ""),          # символ валюты в цене, поле валюты пусто -> CNY
    ("item2", "4", "", "USD"),              # unit_price пуст, есть total ниже -> досчёт
    ("item3", "2", "по запросу", ""),       # не распознано -> предупреждение
    ("item4", "3", "5", "USD"),
]
TOTALS = {"item2": "1 234,56"}  # неразрывный пробел как разделитель тысяч


def fake_chat(messages, **kw):
    t = messages[-1]["content"]
    if "Разбери инвойс" in t:
        out = []
        for i, (n, q, p, c) in enumerate(ROWS):
            out.append({"name_original": n, "quantity": q, "unit_price": p,
                        "total_price": TOTALS.get(n, ""), "currency": c, "row": i + 2})
        return json.dumps(out)
    names = re.findall(r"наименование: (\S+)", t)
    return json.dumps([{"name_ru": "RU " + n, "brand": "B", "material": "m"} for n in names])


gemini.chat, gemini.API_KEY = fake_chat, "test"
wb = openpyxl.Workbook(); wb.active.append(["x"]); b = io.BytesIO(); wb.save(b)
c = TestClient(appmod.app, base_url="https://testserver")
c.post("/api/login", json={"username": "boss", "password": "bosspass123"})
jid = c.post("/api/process", files=[("files", ("inv.xlsx", b.getvalue()))]).json()["job_id"]
for _ in range(50):
    st = c.get(f"/api/status/{jid}").json()
    if st["status"] in ("done", "error"):
        break
    time.sleep(0.1)
print("  статус:", st["status"], "|", st["message"])
for w in st["warnings"]:
    print("  ⚠", w)
check("задача завершена", st["status"] == "done")
ws = openpyxl.load_workbook(io.BytesIO(c.get(f"/api/result/{jid}").content))["Result"]
got = {ws.cell(r, 1).value.replace("RU ", ""): ws.cell(r, 11).value for r in range(3, 3 + len(ROWS))}
check(f"item1: ¥12.5 × 10 = $17.86 (курс 1/7): {got['item1']}", got["item1"] == round(12.5 * 10 / 7, 2))
check(f"item2: total «1 234,56» (нераз. пробел) распознан как $1234.56: {got['item2']}", got["item2"] == 1234.56)
check(f"item4: явная валюта USD, total досчитан = 15: {got['item4']}", got["item4"] == 15.0)
check("предупреждение о нераспознанной цене item3 (товар № 3)",
      any("№ 3" in w and "«по запросу»" in w and "не распознана" in w for w in st["warnings"]))

print("\nИТОГ:", "ВСЁ ОК" if ok else "ЕСТЬ ОШИБКИ")
