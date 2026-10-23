"""Скрытые листы Excel не читаются как источник товаров (кроме случая, когда скрыто всё)."""
import io

import openpyxl

import readers


def _book(states):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for i, st in enumerate(states):
        ws = wb.create_sheet(f"S{i}")
        ws["A1"] = f"товар {i}"
        ws.sheet_state = st
    if "visible" not in states:
        wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_hidden_sheet_skipped():
    assert [s["name"] for s in readers.read_sheets(_book(["hidden", "visible"]))] == ["S1"]


def test_all_visible_read_all():
    assert len(readers.read_sheets(_book(["visible", "visible"]))) == 2


def test_all_hidden_reads_all():
    # Excel не даёт сохранить книгу без видимых листов — проверяем сам отбор
    assert readers._visible(["a", "b"], lambda s: False) == ["a", "b"]
